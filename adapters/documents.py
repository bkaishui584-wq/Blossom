"""文档解析的第二档：装了 MarkItDown 就用它把常见格式读成 markdown。

这一层只负责「本地文件 -> 文本」，并如实说明是哪一档读出来的：
容器格式（docx / xlsx / pptx / pdf）优先交给 MarkItDown，读不出来再退回
tools/workspace.py 自带的朴素解析；纯文本格式根本用不上它。

只读本地文件，不联网、不执行文档里的任何脚本或宏；文件太大直接拒绝，不硬啃。
MarkItDown 没装也不影响任何功能，只是容器格式退回自带解析。
"""

import importlib.util
import os

# 自带解析也能读的容器格式：MarkItDown 读得更整齐（保留标题、列表、表格）。
CONTAINER_SUFFIXES = (".docx", ".xlsx", ".pptx", ".pdf")
# 只有 MarkItDown 认识的格式：没装就只能如实说读不了。
EXTRA_SUFFIXES = (".html", ".htm", ".epub", ".rtf", ".ipynb", ".msg", ".xml")
# 超过这个大小就不硬啃，直接告诉调用方文件太大。
MAX_BYTES = 25 * 1024 * 1024

_engine = None
_failed = ""


def available():
    """包在不在。只看包，不加载。"""
    try:
        return importlib.util.find_spec("markitdown") is not None
    except Exception:
        return False


def too_big(path):
    try:
        return os.path.getsize(path) > MAX_BYTES
    except OSError:
        return False


def failure_reason():
    return _failed


def _load():
    """MarkItDown 实例只建一次；建不起来就记住原因，不反复重试。"""
    global _engine, _failed
    if _engine is not None or _failed:
        return _engine
    try:
        from markitdown import MarkItDown

        # 插件默认关掉：不因为装了别的包就悄悄多出转换器。
        _engine = MarkItDown(enable_plugins=False)
    except Exception as exc:
        _failed = str(exc)
    return _engine


def to_markdown(path, max_chars=6000):
    """把文件读成 markdown 文本。读不出来返回空串，不抛异常。"""
    if not available() or too_big(path):
        return ""
    engine = _load()
    if engine is None:
        return ""
    try:
        result = engine.convert_local(str(path))
    except Exception:
        return ""
    text = str(getattr(result, "text_content", "") or "")
    return text[: int(max_chars or 6000)]


def local_fallback(path, max_chars=6000):
    """MarkItDown 用不上时的第二选择：只用包里本来就有的小库。

    markitdown 的 xlsx 转换器硬依赖 pandas，为了一个格式把 pandas（约 59 MB）塞进安装包
    不划算；openpyxl 只有 2 MB，读出来的表格一样能用。pdf 同理，pdfminer 比自带的朴素抽取准得多。
    读不出来就返回空串，由调用方继续往下退。
    """
    if too_big(path):
        return "", ""
    suffix = os.path.splitext(str(path))[1].lower()
    if suffix == ".xlsx":
        return _xlsx_text(path, max_chars), "openpyxl"
    if suffix == ".pdf":
        return _pdf_text(path, max_chars), "pdfminer"
    return "", ""


def _xlsx_text(path, max_chars=6000):
    """用 openpyxl 逐行读单元格，拼成 a | b | c 这种一行一行的表格文本。"""
    try:
        from openpyxl import load_workbook
    except Exception:
        return ""
    try:
        book = load_workbook(str(path), read_only=True, data_only=True)
    except Exception:
        return ""
    limit = int(max_chars or 6000)
    lines = []
    try:
        for sheet in book.worksheets:
            lines.append("## %s" % (sheet.title or "sheet"))
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if cell is None else str(cell) for cell in row]
                while cells and not cells[-1]:
                    cells.pop()
                if any(cell.strip() for cell in cells):
                    lines.append(" | ".join(cells))
                if sum(len(item) + 1 for item in lines) > limit:
                    break
            if sum(len(item) + 1 for item in lines) > limit:
                break
    except Exception:
        pass
    finally:
        try:
            book.close()
        except Exception:
            pass
    return "\n".join(lines)[:limit]


def _pdf_text(path, max_chars=6000):
    """pdfminer 抽正文。装不上或抽不出都返回空串。"""
    try:
        from pdfminer.high_level import extract_text
    except Exception:
        return ""
    try:
        text = extract_text(str(path))
    except Exception:
        return ""
    return str(text or "")[: int(max_chars or 6000)]


def self_test():
    """真转两份内存里的文件，确认打包之后转换器真的能用。返回 (ok, 说明)。"""
    if not available():
        return False, "没装 MarkItDown"
    engine = _load()
    if engine is None:
        return False, "装了但起不来：%s" % (_failed or "未知原因")
    note = "转换器可用"
    try:
        import io

        html = engine.convert_stream(io.BytesIO(b"<h1>ok</h1>"), file_extension=".html")
        if "ok" not in str(getattr(html, "text_content", "") or ""):
            return False, "html 转换结果不对"
    except Exception as exc:
        return False, "html 转换失败：%s" % exc
    try:
        import io

        from openpyxl import Workbook

        book = Workbook()
        book.active.append(["a", "b"])
        buffer = io.BytesIO()
        book.save(buffer)
        buffer.seek(0)
        sheet = engine.convert_stream(buffer, file_extension=".xlsx")
        if "a" not in str(getattr(sheet, "text_content", "") or ""):
            note = "转换器可用（xlsx 缺 pandas，靠 openpyxl 兜底）"
        else:
            note = "转换器可用（html / xlsx 都过）"
    except ImportError:
        pass
    except Exception:
        note = "转换器可用（xlsx 缺依赖，靠 openpyxl 兜底）"
    return True, note


def describe():
    """给能力页用：装没装、能不能用、为什么不能用。"""
    if not available():
        return {
            "installed": False,
            "ready": False,
            "note": "没装 MarkItDown，容器格式走程序自带的解析",
        }
    if _load() is None:
        return {
            "installed": True,
            "ready": False,
            "note": "装了但起不来：%s" % (_failed or "未知原因"),
        }
    return {"installed": True, "ready": True, "note": "装了，容器格式优先用它"}
